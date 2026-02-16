import secrets
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any, TypeVar

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import and_, case, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import contains_eager, selectinload

from backend.src.auth.service import mask_sensitive_data
from backend.src.common.constants import PaginatedResponse
from backend.src.common.enums import (
    DisputeStatus,
    OrderStatus,
    TransactionType,
    UserRole,
    UserStatus,
)
from backend.src.common.utils.paginaters import get_paginated_list
from backend.src.core.config import settings
from backend.src.core.logging import get_logger
from backend.src.couriers.schemas import CourierCardResponse
from backend.src.disputes.schemas import DisputeCardResponse
from backend.src.models.courier import Courier
from backend.src.models.dispute import Dispute
from backend.src.models.order import Order
from backend.src.models.registration_code import RegistrationCode
from backend.src.models.shop import Shop
from backend.src.models.transaction import Transaction
from backend.src.models.user import User
from backend.src.orders.schemas import OrderCardResponse
from backend.src.shops.schemas import ShopCardResponse

from .schemas import RegistrationCodeResponse

logger = get_logger(__name__)

T = TypeVar("T")


# Константы для улучшения читаемости и поддерживаемости
MAX_GENERATION_ATTEMPTS = 10


def _generate_registration_code() -> str:
    """
    Генерирует случайный код регистрации.
    """
    _validate_code_settings()
    code_characters = settings.auth.code_characters
    code_length = settings.auth.code_length
    return "".join(secrets.choice(code_characters) for _ in range(code_length))


def _validate_code_settings() -> None:
    """
    Проверяет корректность настроек для генерации кода.

    Raises:
        ValueError: если настройки некорректны.
    """
    if not settings.auth.code_characters:
        raise ValueError("Набор символов для генерации кода не может быть пустым.")
    if settings.auth.code_length <= 0:
        raise ValueError("Длина кода должна быть положительной.")


def _create_code_object(
    code: str,
    role: UserRole,
    created_by_admin_id: int,
    now: datetime,
    expires_at: datetime,
) -> RegistrationCode:
    """
    Создает объект RegistrationCode с заданными параметрами.
    """
    return RegistrationCode(
        code=code,
        role=role,
        is_used=False,
        created_by_admin_id=created_by_admin_id,
        created_at=now,
        expires_at=expires_at,
    )


async def _try_save_code(db: AsyncSession, reg_code: RegistrationCode, attempt: int) -> bool:
    """
    Пытается сохранить код в БД.

    Args:
        db: Сессия базы данных.
        reg_code: Объект кода для сохранения.
        attempt: Номер текущей попытки.

    Returns:
        True, если код успешно сохранен, False при коллизии.
    """
    db.add(reg_code)
    try:
        await db.flush()
        await db.refresh(reg_code)
        logger.info(
            f"Сгенерирован уникальный код регистрации '{reg_code.code}' "
            f"для роли '{reg_code.role.value}' на попытке {attempt}."
        )
        return True
    except IntegrityError:
        masked_code = mask_sensitive_data(reg_code.code)
        logger.warning(
            f"Коллизия при генерации кода '{masked_code}' на попытке {attempt}. "
            f"Генерируем новый код."
        )
        await db.rollback()
        return False


def _calculate_expiration_time() -> tuple[datetime, datetime]:
    """
    Вычисляет время создания и истечения кода.

    Returns:
        Кортеж (created_at, expires_at).
    """
    now = datetime.now(UTC)
    expires_at = now + timedelta(hours=settings.auth.registration_code_lifetime_hours)
    return now, expires_at


async def create_registration_code(
    db: AsyncSession, role: UserRole, created_by_admin_id: int
) -> RegistrationCode:
    """
    Генерирует и сохраняет в БД одноразовый код регистрации для пользователя.
    """
    now, expires_at = _calculate_expiration_time()

    try:
        new_code = _generate_registration_code()
    except ValueError as e:
        logger.error(f"Некорректные настройки генерации кода: {e}")
        raise

    for attempt in range(1, MAX_GENERATION_ATTEMPTS + 1):
        reg_code = _create_code_object(new_code, role, created_by_admin_id, now, expires_at)

        if await _try_save_code(db, reg_code, attempt):
            return reg_code

        if attempt < MAX_GENERATION_ATTEMPTS:
            try:
                new_code = _generate_registration_code()
            except ValueError as e:
                logger.error(f"Некорректные настройки генерации кода: {e}")
                raise

    logger.error(
        f"Не удалось сгенерировать уникальный код регистрации после "
        f"{MAX_GENERATION_ATTEMPTS} попыток."
    )
    raise RuntimeError("Не удалось сгенерировать уникальный код регистрации.")


async def get_registration_codes(
    db: AsyncSession,
    page: int,
    limit: int,
    role: UserRole | None = None,
    is_used: bool | None = None,
) -> PaginatedResponse[RegistrationCodeResponse]:
    """
    Получить список кодов регистрации с пагинацией и фильтрацией.
    """
    query = select(RegistrationCode)

    if role:
        query = query.where(RegistrationCode.role == role)

    if is_used is not None:
        query = query.where(RegistrationCode.is_used == is_used)

    # Используем чистый запрос для подсчета без лишних данных
    count_query = select(func.count()).select_from(RegistrationCode)
    if role:
        count_query = count_query.where(RegistrationCode.role == role)
    if is_used is not None:
        count_query = count_query.where(RegistrationCode.is_used == is_used)
    total_count = await db.scalar(count_query)

    query = (
        query.order_by(RegistrationCode.created_at.desc()).offset((page - 1) * limit).limit(limit)
    )
    result = await db.execute(query)
    codes = result.scalars().all()

    response_items = [RegistrationCodeResponse.model_validate(code) for code in codes]
    return PaginatedResponse(total=total_count or 0, items=response_items)


async def get_registration_code_by_id(db: AsyncSession, code_id: int) -> RegistrationCode | None:
    """
    Получить код регистрации по ID.
    """
    return await db.get(RegistrationCode, code_id)


async def deactivate_registration_code(db: AsyncSession, code_id: int) -> RegistrationCode:
    """
    Деактивирует код регистрации (помечает как использованный или просроченный).
    В текущей реализации мы просто помечаем его как использованный, но без привязки к пользователю,
    либо можно установить expires_at в прошлое.
    По требованию: "Деактивировать просрочен = true (менять expires at или is_used хз первый вариант больше нравится)"
    Давайте менять expires_at на текущее время, чтобы он стал просроченным.
    """
    code = await get_registration_code_by_id(db, code_id)
    if not code:
        raise ValueError(f"Код с ID {code_id} не найден.")

    if code.is_used:
        raise ValueError("Код уже использован.")

    # Делаем код просроченным
    code.expires_at = datetime.now(UTC)
    db.add(code)
    await db.commit()
    await db.refresh(code)
    return code


async def get_unused_registration_codes_count(db: AsyncSession) -> dict[Any, int]:
    """
    Получить количество неиспользованных кодов по ролям.
    """
    result = await db.execute(
        select(RegistrationCode.role, func.count())
        .where(RegistrationCode.is_used.is_(False))
        .group_by(RegistrationCode.role)
    )

    counts: dict[UserRole, int] = {UserRole(role): count for role, count in result.all()}
    return {
        UserRole.SHOP: counts.get(UserRole.SHOP, 0),
        UserRole.COURIER: counts.get(UserRole.COURIER, 0),
        UserRole.ADMIN: counts.get(UserRole.ADMIN, 0),
        "total": sum(counts.values()),
    }


# Кэш для системной статистики
_stats_cache = {"data": None, "expires_at": datetime.min.replace(tzinfo=UTC)}


async def get_system_stats(db: AsyncSession) -> dict:
    """
    Получить системную статистику для администраторов.
    Включает счетчики пользователей, заказов и споров.
    Использует кэширование на 30 секунд.
    """
    global _stats_cache
    now = datetime.now(UTC)

    if _stats_cache["data"] and _stats_cache["expires_at"] > now:
        logger.debug("Используем кэшированную системную статистику")
        return _stats_cache["data"]

    today_start = datetime.now(UTC).replace(hour=0, minute=0, second=0, microsecond=0)

    # Статистика в один проход по таблице заказов
    order_stats_stmt = select(
        Order.status,
        func.count(Order.id).label("count"),
        func.sum(case((Order.created_at >= today_start, 1), else_=0)).label("today_count"),
    ).group_by(Order.status)

    order_counts_res = await db.execute(order_stats_stmt)
    orders_data = order_counts_res.all()

    orders = {row[0].value: row[1] for row in orders_data}
    orders_today_count = sum(row[2] or 0 for row in orders_data)

    total_orders = sum(orders.values())
    completed_orders = orders.get(OrderStatus.COMPLETED.value, 0)
    cancelled_orders = orders.get(OrderStatus.CANCELED.value, 0)
    active_orders = total_orders - completed_orders - cancelled_orders

    # Статистика пользователей и курьеров одним запросом
    user_stats_stmt = select(User.role, User.status, func.count(User.id)).group_by(
        User.role, User.status
    )

    user_counts_res = await db.execute(user_stats_stmt)
    users_data = user_counts_res.all()

    users = {}
    active_couriers_count = 0
    total_users = 0

    for role, status, count in users_data:
        users[role.value] = users.get(role.value, 0) + count
        total_users += count
        if role == UserRole.COURIER and status == UserStatus.ACTIVE:
            active_couriers_count += count

    # Статистика споров
    dispute_counts = await db.execute(
        select(Dispute.status, func.count(Dispute.id)).group_by(Dispute.status)
    )
    disputes = {status.value: count for status, count in dispute_counts.all()}
    total_disputes = sum(disputes.values())
    unresolved_disputes = total_disputes - disputes.get(DisputeStatus.RESOLVED.value, 0)

    stats_data = {
        "total_users": total_users,
        "total_admins": users.get(UserRole.ADMIN.value, 0),
        "total_shops": users.get(UserRole.SHOP.value, 0),
        "total_couriers": users.get(UserRole.COURIER.value, 0),
        "total_orders": total_orders,
        "orders_by_status": orders,
        "completed_orders": completed_orders,
        "cancelled_orders": cancelled_orders,
        "active_orders": active_orders,
        "total_disputes": total_disputes,
        "unresolved_disputes": unresolved_disputes,
        "orders_today": orders_today_count or 0,
        "active_couriers": active_couriers_count or 0,
    }

    _stats_cache = {
        "data": stats_data,
        "expires_at": now + timedelta(seconds=30),
    }

    return stats_data


async def get_order_by_id(db: AsyncSession, order_id: int) -> Order | None:
    """
    Получить заказ по ID с полной загрузкой связанных сущностей.
    """
    stmt = (
        select(Order)
        .where(Order.id == order_id)
        .options(
            selectinload(Order.shop),
            selectinload(Order.courier),
            selectinload(Order.history),
            selectinload(Order.dispute),
            selectinload(Order.rating),
        )
    )
    result = await db.execute(stmt)
    return result.scalars().first()


ModelType = TypeVar("ModelType")


async def get_all_couriers(
    db: AsyncSession,
    page: int,
    limit: int,
    status: UserStatus | None,
    search: str | None,
) -> PaginatedResponse[CourierCardResponse]:
    """
    Получает пагинированный список курьеров с возможностью фильтрации и поиска.
    """
    total, couriers = await get_paginated_list(
        db=db,
        model=Courier,
        page=page,
        limit=limit,
        status=status,
        status_field="status",
        status_model=User,
        search=search,
        search_fields=[Courier.full_name, User.username],
        joins=[(User, Courier.user_id == User.id)],
        eager_load_options=[contains_eager(Courier.user)],
        sort_by_field="id",
        sort_desc=False,
    )

    response_items = [CourierCardResponse.model_validate(courier) for courier in couriers]
    return PaginatedResponse(total=total, items=response_items)


async def get_all_shops(
    db: AsyncSession,
    page: int,
    limit: int,
    status: UserStatus | None,
    search: str | None,
) -> PaginatedResponse[ShopCardResponse]:
    """
    Получает пагинированный список магазинов с возможностью фильтрации и поиска.
    """
    total, shops = await get_paginated_list(
        db=db,
        model=Shop,
        page=page,
        limit=limit,
        status=status,
        status_field="status",
        status_model=User,
        search=search,
        search_fields=[Shop.name, User.username],
        joins=[(User, Shop.user_id == User.id)],
        eager_load_options=[contains_eager(Shop.user)],
        sort_by_field="id",
        sort_desc=False,
    )

    response_items = [ShopCardResponse.model_validate(shop) for shop in shops]
    return PaginatedResponse(total=total, items=response_items)


async def get_all_orders(
    db: AsyncSession,
    page: int,
    limit: int,
    status: OrderStatus | None,
    search: str | None,
    current: bool | None = None,
) -> PaginatedResponse[OrderCardResponse]:
    """
    Получает пагинированный список заказов с возможностью фильтрации и поиска.
    """
    search_fields = [
        Shop.name,
        Courier.full_name,
        Order.description,
    ]

    additional_filters = []
    if current is True:
        additional_filters.append(Order.status.in_(OrderStatus.active_statuses()))
    elif current is False:
        additional_filters.append(Order.status.in_(OrderStatus.completed_statuses()))

    total, orders = await get_paginated_list(
        db=db,
        model=Order,
        page=page,
        limit=limit,
        status=status,
        status_field="status",
        search=search,
        search_fields=search_fields,
        joins=[
            (Shop, Order.shop_id == Shop.id, True),
            (Courier, Order.courier_id == Courier.id, True),
        ],
        eager_load_options=[
            contains_eager(Order.shop),
            contains_eager(Order.courier),
        ],
        sort_by_field="created_at",
        sort_desc=True,
        additional_filters=additional_filters,
    )

    response_items = [OrderCardResponse.model_validate(order) for order in orders]
    return PaginatedResponse(total=total, items=response_items)


async def get_all_disputes(
    db: AsyncSession,
    page: int,
    limit: int,
    status: DisputeStatus | None,
    search: str | None,
) -> PaginatedResponse[DisputeCardResponse]:
    """
    Получить все споры с пагинацией и фильтрацией.
    Доступно только администраторам.
    """
    search_fields = [
        Dispute.description,
        Shop.name,
        Courier.full_name,
        User.username,
    ]

    total, disputes = await get_paginated_list(
        db=db,
        model=Dispute,
        page=page,
        limit=limit,
        status=status,
        status_field="status",
        search=search,
        search_fields=search_fields,
        joins=[
            (Order, Dispute.order_id == Order.id),
            (Shop, Order.shop_id == Shop.id, True),  # type: ignore
            (Courier, Order.courier_id == Courier.id, True),  # type: ignore
            (User, Dispute.opened_by_user_id == User.id),
        ],
        eager_load_options=[
            contains_eager(Dispute.order).contains_eager(Order.shop),
            contains_eager(Dispute.order).contains_eager(Order.courier),
            contains_eager(Dispute.opened_by_user),
        ],
        sort_by_field="created_at",
        sort_desc=True,
    )

    response_items = [DisputeCardResponse.from_dispute(d) for d in disputes]

    return PaginatedResponse(total=total, items=response_items)


async def _get_last_cash_collection_date(db: AsyncSession, shop_id: int) -> datetime | None:
    """
    Получить дату последней инкассации (CASH_COLLECTION) для магазина.
    Если инкассаций не было, вернуть дату первого заказа.
    """
    # Ищем последнюю CASH_COLLECTION
    stmt = (
        select(Transaction.created_at)
        .join(Shop, Shop.user_id == Transaction.user_id)
        .where(
            and_(
                Shop.id == shop_id,
                Transaction.type == TransactionType.CASH_COLLECTION,
            )
        )
        .order_by(Transaction.created_at.desc())
        .limit(1)
    )
    result = await db.execute(stmt)
    last_collection = result.scalar_one_or_none()

    if last_collection:
        return last_collection

    # Если нет инкассаций, берём дату первого заказа
    stmt = (
        select(Order.created_at)
        .where(Order.shop_id == shop_id)
        .order_by(Order.created_at.asc())
        .limit(1)
    )
    result = await db.execute(stmt)
    first_order = result.scalar_one_or_none()

    return first_order


async def _get_last_payout_date(db: AsyncSession, courier_id: int) -> datetime | None:
    """
    Получить дату последней выплаты (PAYOUT) для курьера.
    Если выплат не было, вернуть дату первой транзакции ORDER_CREDIT.
    """
    # Ищем последнюю PAYOUT
    stmt = (
        select(Transaction.created_at)
        .join(Courier, Courier.user_id == Transaction.user_id)
        .where(
            and_(
                Courier.id == courier_id,
                Transaction.type == TransactionType.PAYOUT,
            )
        )
        .order_by(Transaction.created_at.desc())
        .limit(1)
    )
    result = await db.execute(stmt)
    last_payout = result.scalar_one_or_none()

    if last_payout:
        return last_payout

    # Если нет выплат, берём первую ORDER_CREDIT
    stmt = (
        select(Transaction.created_at)
        .join(Courier, Courier.user_id == Transaction.user_id)
        .where(
            and_(
                Courier.id == courier_id,
                Transaction.type == TransactionType.ORDER_CREDIT,
            )
        )
        .order_by(Transaction.created_at.asc())
        .limit(1)
    )
    result = await db.execute(stmt)
    first_credit = result.scalar_one_or_none()

    return first_credit


def _create_excel_response(workbook: Workbook, filename: str) -> StreamingResponse:
    """
    Создаёт StreamingResponse с Excel файлом.
    """
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

    return StreamingResponse(output, headers=headers, media_type=headers["Content-Type"])


def _format_header_row(ws, headers: list[str]):
    """
    Форматирует заголовок таблицы (жирный шрифт, выравнивание по центру).
    """
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _add_total_row(ws, row_num: int, totals: dict[str, any]):
    """
    Добавляет итоговую строку с суммами.

    Args:
        ws: Worksheet
        row_num: Номер строки для итогов
        totals: Словарь {column_letter: value}, например {"A": "ИТОГО:", "G": 15000}
    """
    for col_letter, value in totals.items():
        cell = ws[f"{col_letter}{row_num}"]
        cell.value = value
        cell.font = Font(bold=True)


async def export_shop_statistics(
    db: AsyncSession,
    shop_id: int,
    date_from: datetime,
    date_to: datetime,
    stats_type: str,
    from_last_payment: bool = False,
) -> StreamingResponse:
    """
    Экспорт статистики по конкретному магазину в Excel.

    Args:
        db: Сессия базы данных
        shop_id: ID магазина
        date_from: Начальная дата (игнорируется если from_last_payment=True)
        date_to: Конечная дата
        stats_type: "common" или "advanced"
        from_last_payment: Если True, берётся дата последней CASH_COLLECTION
    """
    # Проверяем существование магазина
    shop = await db.get(Shop, shop_id)
    if not shop:
        raise ValueError(f"Магазин с ID {shop_id} не найден")

    # Определяем начальную дату
    if from_last_payment:
        last_payment = await _get_last_cash_collection_date(db, shop_id)
        if last_payment:
            date_from = last_payment
        else:
            logger.warning(f"Для магазина ID {shop_id} нет ни платежей, ни заказов")

    # Получаем заказы
    stmt = (
        select(Order)
        .options(
            selectinload(Order.courier),  # ← Загружаем курьера
            selectinload(Order.dispute),  # ← Загружаем спор
            selectinload(Order.rating),  # ← Загружаем рейтинг
            selectinload(Order.shop),  # ← Загружаем магазин
        )
        .where(
            and_(
                Order.shop_id == shop_id,
                Order.created_at >= date_from,
                Order.created_at <= date_to,
            )
        )
        .order_by(Order.created_at.desc())
    )
    result = await db.execute(stmt)
    orders = result.scalars().all()

    # Создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика заказов"

    # Заголовки
    if stats_type == "common":
        headers = [
            "ID заказа",
            "Дата создания",
            "Дата завершения",
            "Тип заказа",
            "Описание",
            "Курьер",
            "Стоимость",
            "Спор",
            "Штраф",
        ]
    else:  # advanced
        headers = [
            "ID заказа",
            "Дата создания",
            "Дата завершения",
            "Длительность",
            "Тип заказа",
            "Описание",
            "Курьер",
            "Рейтинг",
            "Стоимость",
            "Оплата курьеру",
            "Комиссия",
            "Штраф",
            "Спор",
        ]

    _format_header_row(ws, headers)

    # Заполняем данные
    total_price = 0
    total_courier_payment = 0
    total_service_fee = 0
    total_fee = 0

    for idx, order in enumerate(orders, start=2):
        # Получаем связанные данные
        courier_name = order.courier.full_name if order.courier else "Не назначен"
        dispute_text = "Да" if order.dispute else "Нет"

        # Штраф из dispute
        fee = 0
        if order.dispute and order.dispute.fine_amount:
            # Если оштрафован курьер - это плюс для магазина
            if order.dispute.fined_user_id == order.courier_id:
                fee = float(order.dispute.fine_amount)
            # Если оштрафован магазин - это минус
            elif order.dispute.fined_user_id == shop.user_id:
                fee = -float(order.dispute.fine_amount)

        total_fee += fee

        if stats_type == "common":
            ws.cell(row=idx, column=1, value=order.id)
            ws.cell(row=idx, column=2, value=order.created_at.strftime("%Y-%m-%d %H:%M"))
            ws.cell(
                row=idx,
                column=3,
                value=order.completed_at.strftime("%Y-%m-%d %H:%M") if order.completed_at else "",
            )
            ws.cell(row=idx, column=4, value=order.order_type.value)
            ws.cell(row=idx, column=5, value=order.description or "")
            ws.cell(row=idx, column=6, value=courier_name)
            ws.cell(row=idx, column=7, value=float(order.price))
            ws.cell(row=idx, column=8, value=dispute_text)
            ws.cell(row=idx, column=9, value=fee)

            total_price += float(order.price)

        else:  # advanced
            # Получаем рейтинг
            rating = order.rating.rating if order.rating else None

            # Вычисляем courier_payment и service_fee
            # Предполагаем, что это хранится в транзакциях
            courier_payment = 0
            service_fee = 0

            # Находим транзакции по заказу
            transactions_stmt = select(Transaction).where(Transaction.order_id == order.id)
            trans_result = await db.execute(transactions_stmt)
            transactions = trans_result.scalars().all()

            for trans in transactions:
                if trans.type == TransactionType.ORDER_CREDIT:
                    courier_payment = float(trans.amount)
                elif trans.type == TransactionType.SERVICE_FEE:
                    service_fee = float(trans.amount)

            # Длительность
            duration = ""
            if order.completed_at:
                delta = order.completed_at - order.created_at
                hours = delta.total_seconds() / 3600
                duration = f"{hours:.1f}ч"

            ws.cell(row=idx, column=1, value=order.id)
            ws.cell(row=idx, column=2, value=order.created_at.strftime("%Y-%m-%d %H:%M"))
            ws.cell(
                row=idx,
                column=3,
                value=order.completed_at.strftime("%Y-%m-%d %H:%M") if order.completed_at else "",
            )
            ws.cell(row=idx, column=4, value=duration)
            ws.cell(row=idx, column=5, value=order.order_type.value)
            ws.cell(row=idx, column=6, value=order.description or "")
            ws.cell(row=idx, column=7, value=courier_name)
            ws.cell(row=idx, column=8, value=rating if rating else "")
            ws.cell(row=idx, column=9, value=float(order.price))
            ws.cell(row=idx, column=10, value=courier_payment)
            ws.cell(row=idx, column=11, value=service_fee)
            ws.cell(row=idx, column=12, value=fee)
            ws.cell(row=idx, column=13, value=dispute_text)

            total_price += float(order.price)
            total_courier_payment += courier_payment
            total_service_fee += service_fee

    # Добавляем итоговую строку
    total_row = len(orders) + 2

    if stats_type == "common":
        _add_total_row(
            ws,
            total_row,
            {
                "F": "ИТОГО:",
                "G": total_price,
                "I": total_fee,
            },
        )
    else:  # advanced
        _add_total_row(
            ws,
            total_row,
            {
                "H": "ИТОГО:",
                "I": total_price,
                "J": total_courier_payment,
                "K": total_service_fee,
                "L": total_fee,
            },
        )

    # Формируем имя файла
    filename = f"shop_{shop_id}_{stats_type}_{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}.xlsx"

    return _create_excel_response(wb, filename)


async def export_courier_statistics(
    db: AsyncSession,
    courier_id: int,
    date_from: datetime,
    date_to: datetime,
    stats_type: str,
    from_last_payout: bool = False,
) -> StreamingResponse:
    """
    Экспорт статистики по конкретному курьеру в Excel.
    """
    # Проверяем существование курьера
    courier = await db.get(Courier, courier_id)
    if not courier:
        raise ValueError(f"Курьер с ID {courier_id} не найден")

    # Определяем начальную дату
    if from_last_payout:
        last_payout = await _get_last_payout_date(db, courier_id)
        if last_payout:
            date_from = last_payout
        else:
            logger.warning(f"Для курьера ID {courier_id} нет ни выплат, ни заработков")

    # Получаем заказы
    stmt = (
        select(Order)
        .options(
            selectinload(Order.courier),
            selectinload(Order.dispute),
            selectinload(Order.rating),
            selectinload(Order.shop),
        )
        .where(
            and_(
                Order.courier_id == courier_id,
                Order.created_at >= date_from,
                Order.created_at <= date_to,
            )
        )
        .order_by(Order.created_at.desc())
    )
    result = await db.execute(stmt)
    orders = result.scalars().all()

    # Создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика заказов"

    # Заголовки
    if stats_type == "common":
        headers = [
            "ID заказа",
            "Дата создания",
            "Дата завершения",
            "Длительность",
            "Магазин",
            "Тип заказа",
            "Рейтинг",
            "Комментарий",
            "Заработок",
            "Штраф",
            "Комиссия",
            "Спор",
            "Результат спора",
        ]
    else:  # advanced
        headers = [
            "ID заказа",
            "Дата создания",
            "Дата завершения",
            "Длительность",
            "Магазин",
            "Тип заказа",
            "Рейтинг",
            "Комментарий",
            "Заработок",
            "Стоимость",
            "Комиссия",
            "Штраф",
            "Спор",
            "Результат спора",
        ]

    _format_header_row(ws, headers)

    # Заполняем данные
    total_earnings = 0
    total_price = 0
    total_service_fee = 0
    total_fee = 0

    for idx, order in enumerate(orders, start=2):
        shop_name = order.shop.name if order.shop else "Неизвестно"
        rating = order.rating.rating if order.rating else ""
        rating_comment = order.rating.comment if order.rating else ""

        dispute_text = "Да" if order.dispute else "Нет"
        dispute_result = ""
        if order.dispute and order.dispute.resolution_type:
            dispute_result = order.dispute.resolution_type.value

        # Штраф
        fee = 0
        if order.dispute and order.dispute.fine_amount:
            # Если оштрафован курьер - это минус
            if order.dispute.fined_user_id == courier.user_id:
                fee = -float(order.dispute.fine_amount)
            # Если оштрафован магазин - это плюс для курьера
            elif order.dispute.fined_user_id == order.shop.user_id:
                fee = float(order.dispute.fine_amount)

        total_fee += fee

        # Получаем транзакции
        earnings = 0
        service_fee = 0

        transactions_stmt = select(Transaction).where(
            and_(
                Transaction.order_id == order.id,
                Transaction.user_id == courier.user_id,
            )
        )
        trans_result = await db.execute(transactions_stmt)
        transactions = trans_result.scalars().all()

        for trans in transactions:
            if trans.type == TransactionType.ORDER_CREDIT:
                earnings = float(trans.amount)
            elif trans.type == TransactionType.SERVICE_FEE:
                service_fee = float(trans.amount)

        # Длительность
        duration = ""
        if order.completed_at:
            delta = order.completed_at - order.created_at
            hours = delta.total_seconds() / 3600
            duration = f"{hours:.1f}ч"

        if stats_type == "common":
            ws.cell(row=idx, column=1, value=order.id)
            ws.cell(row=idx, column=2, value=order.created_at.strftime("%Y-%m-%d %H:%M"))
            ws.cell(
                row=idx,
                column=3,
                value=order.completed_at.strftime("%Y-%m-%d %H:%M") if order.completed_at else "",
            )
            ws.cell(row=idx, column=4, value=duration)
            ws.cell(row=idx, column=5, value=shop_name)
            ws.cell(row=idx, column=6, value=order.order_type.value)
            ws.cell(row=idx, column=7, value=rating)
            ws.cell(row=idx, column=8, value=rating_comment)
            ws.cell(row=idx, column=9, value=earnings)
            ws.cell(row=idx, column=10, value=fee)
            ws.cell(row=idx, column=11, value=service_fee)
            ws.cell(row=idx, column=12, value=dispute_text)
            ws.cell(row=idx, column=13, value=dispute_result)

            total_earnings += earnings

        else:  # advanced
            ws.cell(row=idx, column=1, value=order.id)
            ws.cell(row=idx, column=2, value=order.created_at.strftime("%Y-%m-%d %H:%M"))
            ws.cell(
                row=idx,
                column=3,
                value=order.completed_at.strftime("%Y-%m-%d %H:%M") if order.completed_at else "",
            )
            ws.cell(row=idx, column=4, value=duration)
            ws.cell(row=idx, column=5, value=shop_name)
            ws.cell(row=idx, column=6, value=order.order_type.value)
            ws.cell(row=idx, column=7, value=rating)
            ws.cell(row=idx, column=8, value=rating_comment)
            ws.cell(row=idx, column=9, value=earnings)
            ws.cell(row=idx, column=10, value=float(order.price))
            ws.cell(row=idx, column=11, value=service_fee)
            ws.cell(row=idx, column=12, value=fee)
            ws.cell(row=idx, column=13, value=dispute_text)
            ws.cell(row=idx, column=14, value=dispute_result)

            total_earnings += earnings
            total_price += float(order.price)
            total_service_fee += service_fee

    # Добавляем итоговую строку
    total_row = len(orders) + 2

    if stats_type == "common":
        _add_total_row(
            ws,
            total_row,
            {
                "H": "ИТОГО:",
                "I": total_earnings,
                "J": total_fee,
            },
        )
    else:  # advanced
        _add_total_row(
            ws,
            total_row,
            {
                "H": "ИТОГО:",
                "I": total_earnings,
                "J": total_price,
                "K": total_service_fee,
                "L": total_fee,
            },
        )

    # Формируем имя файла
    filename = f"courier_{courier_id}_{stats_type}_{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(wb, filename)


async def export_all_shops_statistics(
    db: AsyncSession,
    date_from: datetime,
    date_to: datetime,
) -> StreamingResponse:
    """
    Экспорт статистики по всем заказам всех магазинов в Excel.
    """
    # Получаем все заказы за период
    stmt = (
        select(Order)
        .options(
            selectinload(Order.courier),
            selectinload(Order.dispute),
            selectinload(Order.rating),
            selectinload(Order.shop),
        )
        .where(
            and_(
                Order.created_at >= date_from,
                Order.created_at <= date_to,
            )
        )
        .order_by(Order.created_at.desc())
    )
    result = await db.execute(stmt)
    orders = result.scalars().all()

    # Создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Все заказы"

    # Заголовки
    headers = [
        "ID заказа",
        "Дата создания",
        "Дата завершения",
        "Длительность",
        "Тип заказа",
        "Описание",
        "Курьер",
        "Рейтинг",
        "Стоимость",
        "Оплата курьеру",
        "Комиссия",
        "Штраф",
        "Спор",
    ]

    _format_header_row(ws, headers)

    # Заполняем данные
    total_price = 0
    total_courier_payment = 0
    total_service_fee = 0
    total_fee = 0

    for idx, order in enumerate(orders, start=2):
        courier_name = order.courier.full_name if order.courier else "Не назначен"
        rating = order.rating.rating if order.rating else ""
        dispute_text = "Да" if order.dispute else "Нет"

        # Штраф
        fee = 0
        if order.dispute and order.dispute.fine_amount:
            fee = float(order.dispute.fine_amount)

        # Получаем транзакции
        courier_payment = 0
        service_fee = 0

        transactions_stmt = select(Transaction).where(Transaction.order_id == order.id)
        trans_result = await db.execute(transactions_stmt)
        transactions = trans_result.scalars().all()

        for trans in transactions:
            if trans.type == TransactionType.ORDER_CREDIT:
                courier_payment = float(trans.amount)
            elif trans.type == TransactionType.SERVICE_FEE:
                service_fee = float(trans.amount)

        # Длительность
        duration = ""
        if order.completed_at:
            delta = order.completed_at - order.created_at
            hours = delta.total_seconds() / 3600
            duration = f"{hours:.1f}ч"

        ws.cell(row=idx, column=1, value=order.id)
        ws.cell(row=idx, column=2, value=order.created_at.strftime("%Y-%m-%d %H:%M"))
        ws.cell(
            row=idx,
            column=3,
            value=order.completed_at.strftime("%Y-%m-%d %H:%M") if order.completed_at else "",
        )
        ws.cell(row=idx, column=4, value=duration)
        ws.cell(row=idx, column=5, value=order.order_type.value)
        ws.cell(row=idx, column=6, value=order.description or "")
        ws.cell(row=idx, column=7, value=courier_name)
        ws.cell(row=idx, column=8, value=rating)
        ws.cell(row=idx, column=9, value=float(order.price))
        ws.cell(row=idx, column=10, value=courier_payment)
        ws.cell(row=idx, column=11, value=service_fee)
        ws.cell(row=idx, column=12, value=fee)
        ws.cell(row=idx, column=13, value=dispute_text)

        total_price += float(order.price)
        total_courier_payment += courier_payment
        total_service_fee += service_fee
        total_fee += fee

    # Добавляем итоговую строку
    total_row = len(orders) + 2
    _add_total_row(
        ws,
        total_row,
        {
            "H": "ИТОГО:",
            "I": total_price,
            "J": total_courier_payment,
            "K": total_service_fee,
            "L": total_fee,
        },
    )

    # Формируем имя файла
    filename = f"all_shops_{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(wb, filename)
